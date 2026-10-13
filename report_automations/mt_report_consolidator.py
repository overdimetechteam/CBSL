import pandas as pd
import os
import re
import sys
import logging
import shutil
import argparse
from copy import copy
from calendar import monthrange
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SCRIPT_ROOT = os.path.dirname(os.path.abspath(__file__))

parent_directory = Path(SCRIPT_ROOT).parent

logger = logging.getLogger(__name__)

# Allow passing the <date-folder> explicitly; fallback to auto-detect
parser = argparse.ArgumentParser(description='Process summary using new folder structure')
parser.add_argument('--working-dir', type=str, help='Path to working date folder (working/NBD-QF-23-C3-C10-SF-FI/<date-folder>)')
parser.add_argument('--report-date', type=str, help='Report date in YYYY-MM-DD format (same argument passed to process_summary.py); used to pick the 3-month window and the MT-<month>.xlsx filename')
args, unknown = parser.parse_known_args()

# Resolve the report date the same way process_summary.py does, so that when this
# script is run standalone it uses the same date it would receive when imported/called
# from process_summary.py via consolidate_mt_reports(report_date_obj).
if args.report_date:
    try:
        target_report_date = datetime.strptime(args.report_date, "%Y-%m-%d")
    except ValueError:
        print(f"Warning: could not parse --report-date '{args.report_date}', using current date instead")
        target_report_date = datetime.now()
else:
    target_report_date = datetime.now()

# Last date in completed_dates.txt - names the previous run's MT workbook
# ("MT - <Mon YYYY>.xlsx"). Only known when the folder is auto-detected below.
PREV_COMPLETED_DATE = None

if args.working_dir:
    provided_path = Path(args.working_dir)
    if not provided_path.exists() or not provided_path.is_dir():
        raise FileNotFoundError(f"Working directory not found: {provided_path}")
    # If the provided path is the base folder (working/NBD-QF-23-C3-C10-SF-FI), pick its single <date-folder>
    subdirs = [d for d in provided_path.iterdir() if d.is_dir()]
    if subdirs and not (provided_path / "Input").is_dir():
        if len(subdirs) > 1:
            print(f"Warning: multiple <date-folder>s found under base directory, using the first: {subdirs[0].name}")
        date_folder = subdirs[0]
    else:
        # Treat the provided path as the <date-folder>
        date_folder = provided_path
    output_path = date_folder
else:
    # Source folder: outputs/NBD-QF-23-C3-C10-SF-FI/<last completed date>/ - the
    # UI copies the uploaded Client Available Limits reports into its Input/MT.
    base_working_dir = parent_directory / "outputs" / "NBD-QF-23-C3-C10-SF-FI"
    if not base_working_dir.exists():
        raise FileNotFoundError(f"Working directory not found: {base_working_dir}")

    # Look for completed_dates.txt file
    completed_dates_file = base_working_dir / "completed_dates.txt"
    if not completed_dates_file.exists():
        raise FileNotFoundError(f"completed_dates.txt not found in: {base_working_dir}")

    # Read the date from the text file (get the last line)
    with open(completed_dates_file, 'r') as f:
        lines = [line.strip() for line in f.readlines() if line.strip()]

    if not lines:
        raise ValueError(f"completed_dates.txt is empty in: {base_working_dir}")

    # Get the last date from the file
    target_date = lines[-1]
    PREV_COMPLETED_DATE = target_date
    print(f"Target date from completed_dates.txt: {target_date}")

    # Find all folders that match the target date (including versioned ones)
    date_folders = []
    pattern = re.compile(rf"^{re.escape(target_date)}(\(\d+\))?$")

    for folder in base_working_dir.iterdir():
        if folder.is_dir() and pattern.match(folder.name):
            date_folders.append(folder)

    if not date_folders:
        raise FileNotFoundError(f"No folder found for date '{target_date}' in: {base_working_dir}")

    # Sort folders to get the highest version
    # Folders are sorted by: base date, then version number (0 if no version)
    def get_version_number(folder_name):
        """Extract version number from folder name like '31-07-2025(1)' -> 1, or '31-07-2025' -> 0"""
        match = re.search(r'\((\d+)\)$', folder_name)
        return int(match.group(1)) if match else 0

    # Sort by version number (highest first)
    date_folders.sort(key=lambda f: get_version_number(f.name), reverse=True)

    # Select the folder with the highest version
    date_folder = date_folders[0]
    print(f"Selected date folder (highest version): {date_folder.name}")

    # Store the output path
    output_path = date_folder

print(f"Using date folder: {date_folder}")
print(f"Output path: {output_path}")

# Folder paths (operate strictly within <date-folder>)
INPUT_FOLDER = date_folder / "Input" / "MT"
OUTPUT_FOLDER = date_folder

print(f"Input folder: {INPUT_FOLDER}")
print(f"Output folder: {OUTPUT_FOLDER}")


MT_NO_RE = re.compile(r'^MT\d+$', re.IGNORECASE)
LENDING_HEADER = 'Actual Lending(Rs.)'
MISSING_RECORDS_SHEET = 'MT'


def _month_end(year, month):
    return datetime(year, month, monthrange(year, month)[1])


def _quarter_months(date):
    """The report month and the 2 before it, oldest first, as month-end dates.
    e.g. 30-Sep-2025 -> [31-07-2025, 31-08-2025, 30-09-2025]"""
    months = []
    for i in range(2, -1, -1):
        month = date.month - i
        year = date.year
        if month <= 0:
            month += 12
            year -= 1
        months.append(_month_end(year, month))
    return months


def _month_before(date):
    first = date.replace(day=1) - timedelta(days=1)
    return _month_end(first.year, first.month)


def _to_number(value):
    """Report cells are numbers, but some are text like '(1,234.50)' or '(-0.00)'."""
    if value is None or value == '':
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(',', '')
    negative = text.startswith('(') and text.endswith(')')
    text = text.strip('()')
    try:
        number = float(text)
    except ValueError:
        return 0.0
    return -abs(number) if negative else number


def _find_header(ws, label, max_rows=15):
    """(row, column) of the first cell whose text is exactly `label`."""
    for row in ws.iter_rows(min_row=1, max_row=min(max_rows, ws.max_row)):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == label:
                return cell.row, cell.column
    return None


def read_lending_sheet(ws, source):
    """Read a Client Available Limits layout - an ERP report, or a month sheet
    copied from one - into {MT No.: (Client Name, Actual Lending)}.

    Only rows whose MT No. looks like 'MT0001' are kept, so the title rows,
    header rows, blank rows and the 'Total' row are skipped."""
    mt = _find_header(ws, 'MT No.')
    name = _find_header(ws, 'Client Name')
    lending = _find_header(ws, LENDING_HEADER)
    if not (mt and name and lending):
        raise ValueError(f"{source}: could not find the 'MT No.', 'Client Name' and "
                         f"'{LENDING_HEADER}' headers in the first rows")

    records = {}
    for r in range(mt[0] + 1, ws.max_row + 1):
        value = ws.cell(r, mt[1]).value
        if value is None:
            continue
        mt_no = str(value).strip().upper()
        if not MT_NO_RE.match(mt_no):
            continue
        if mt_no in records:
            logger.warning(f"{source}: duplicate {mt_no} on row {r} - keeping the first row")
            continue
        client = ws.cell(r, name[1]).value
        records[mt_no] = (str(client).strip() if client else '', _to_number(ws.cell(r, lending[1]).value))
    return records


def _copy_report_to_sheet(src_ws, dst_ws):
    """Copy the report table - from its header rows down to the 'Total' row,
    i.e. without the company title lines above or the form code / user name
    footer below - into dst_ws, keeping merged header cells and widths.
    Gives the same layout as the month sheets of earlier MT workbooks."""
    start = _find_header(src_ws, 'MT No.')[0]
    last = start
    for r in range(start, src_ws.max_row + 1):
        first_cell = src_ws.cell(r, 1).value
        if isinstance(first_cell, str) and first_cell.strip().lower() == 'total':
            last = r
            break
        if any(src_ws.cell(r, c).value not in (None, '') for c in range(1, src_ws.max_column + 1)):
            last = r
    offset = start - 1
    for r in range(start, last + 1):
        for c in range(1, src_ws.max_column + 1):
            dst_ws.cell(r - offset, c).value = src_ws.cell(r, c).value
    for merged in src_ws.merged_cells.ranges:
        if merged.min_row >= start and merged.max_row <= last:
            dst_ws.merge_cells(start_row=merged.min_row - offset, start_column=merged.min_col,
                               end_row=merged.max_row - offset, end_column=merged.max_col)
    for letter, dim in src_ws.column_dimensions.items():
        if dim.width:
            dst_ws.column_dimensions[letter].width = dim.width


def _previous_mt_file(months):
    """The previous run's MT workbook: 'MT - <Mon YYYY of the last completed
    date>.xlsx' (e.g. MT - Jul 2025 for a Sep run whose last completed date is
    31-07-2025); falls back to 'MT - <month before the quarter>.xlsx'. A file
    for the report month or later is never used (it would be this run's output)."""
    report_month = months[-1]
    candidates = []
    if PREV_COMPLETED_DATE:
        for fmt in ("%d-%m-%Y", "%d/%m/%Y"):
            try:
                candidates.append(datetime.strptime(PREV_COMPLETED_DATE, fmt))
                break
            except ValueError:
                pass
    candidates.append(_month_before(months[0]))

    tried = []
    for dt in candidates:
        path = INPUT_FOLDER / f"MT - {dt.strftime('%b %Y')}.xlsx"
        if (dt.year, dt.month) >= (report_month.year, report_month.month):
            logger.info(f"Skipping {path.name}: not earlier than the report month")
            continue
        tried.append(str(path))
        if path.exists():
            return path
    raise FileNotFoundError("Previous MT workbook not found. Tried: " + "; ".join(tried))


def _previous_genders(ws):
    """{MT No.: Gender} from the previous workbook's Dis sheet. Handles both the
    layout this script writes (MT No., Client Name, Gender, ...) and the older
    hand-made one (MT No., Gender, Client Name, ... with no header text)."""
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    if 'MT No.' in header and 'Gender' in header:
        mt_col, gender_col = header.index('MT No.') + 1, header.index('Gender') + 1
    else:
        mt_col, gender_col = 1, 2
    genders = {}
    for r in range(2, ws.max_row + 1):
        mt_no = ws.cell(r, mt_col).value
        gender = ws.cell(r, gender_col).value
        if mt_no and gender not in (None, ''):
            genders[str(mt_no).strip().upper()] = str(gender).strip()
    return genders


def _previous_dis_closing(ws):
    """{MT No.: last column value} from the previous Dis sheet - only used as
    the opening balance when the previous workbook has no month sheet."""
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    mt_col = header.index('MT No.') + 1 if 'MT No.' in header else 1
    last_col = max(c for c in range(1, ws.max_column + 1)
                   if any(ws.cell(r, c).value not in (None, '') for r in range(1, min(ws.max_row, 20) + 1)))
    closing = {}
    for r in range(2, ws.max_row + 1):
        mt_no = ws.cell(r, mt_col).value
        if mt_no and MT_NO_RE.match(str(mt_no).strip()):
            closing[str(mt_no).strip().upper()] = _to_number(ws.cell(r, last_col).value)
    return closing


def _mt_sort_key(mt_no):
    digits = re.sub(r'\D', '', mt_no)
    return (int(digits) if digits else 0, mt_no)


def _write_dis_sheet(wb, df):
    if 'Dis' in wb.sheetnames:
        del wb['Dis']
    ws = wb.create_sheet('Dis')
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for col in range(4, len(df.columns) + 1):
        for r in range(2, len(df) + 2):
            ws.cell(r, col).number_format = '#,##0.00'
    ws.column_dimensions['A'].width = 10
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 12
    for col in range(4, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24
    return ws


def _write_disbursement_sheet(wb, df, month_labels, gl=None):
    """Rebuild the Disbursement sheet from the Dis sheet: one row per client
    with the monthly movement in Actual Lending, a Total row, the gender
    summary and the Net MT Disbursement block - all as formulas, so they
    follow the Dis values. `gl` holds the GL 1846 balances (see
    _write_disbursement_summary)."""
    if 'Disbursement' in wb.sheetnames:
        ws = wb['Disbursement']
        for merged in list(ws.merged_cells.ranges):
            if merged.max_row >= 3:
                ws.unmerge_cells(str(merged))
        for r in range(3, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).value = None
    else:
        ws = wb.create_sheet('Disbursement', 0)
        ws['D2'] = 'Disbursement'

    first, last = 4, 3 + len(df)
    for c, title in enumerate(['MT No.', 'Client Name', 'Male/Female/Other'] + month_labels + ['Total'], start=1):
        ws.cell(3, c).value = title

    for i, row in enumerate(df.itertuples(index=False)):
        r, dis_row = first + i, 2 + i   # Disbursement row 4 <-> Dis row 2
        ws.cell(r, 1).value = row[0]
        ws.cell(r, 2).value = row[1]
        ws.cell(r, 3).value = row[2] if isinstance(row[2], str) else None
        ws.cell(r, 4).value = f"=Dis!E{dis_row}-Dis!D{dis_row}"
        ws.cell(r, 5).value = f"=Dis!F{dis_row}-Dis!E{dis_row}"
        ws.cell(r, 6).value = f"=Dis!G{dis_row}-Dis!F{dis_row}"
        ws.cell(r, 7).value = f"=SUM(D{r}:F{r})"

    total_row = max(last, first - 1) + 1
    ws.cell(total_row, 1).value = "Total"
    for c in range(4, 8):
        letter = get_column_letter(c)
        ws.cell(total_row, c).value = f"=SUM({letter}{first}:{letter}{max(last, first)})"
    for r in range(first, total_row + 1):
        for c in range(4, 8):
            ws.cell(r, c).number_format = '#,##0.00'

    _write_disbursement_summary(ws, df, first, max(last, first), gl)
    return ws


SUMMARY_FILL = PatternFill('solid', fgColor='DDEBF7')
SPLIT_FILL = PatternFill('solid', fgColor='FFFF00')
THIN, DOUBLE = Side(style='thin'), Side(style='double')


def _write_disbursement_summary(ws, df, first, last, gl=None):
    """Right-hand side of the Disbursement sheet, laid out like MT - Jul 2025.xlsx.
    The report reads fixed cells of it (K10:K13, L17:L19, N17), so the rows
    below must stay where they are.

      J3:L7   gender summary (was a PivotTable): Female, Male, Corporate, Grand Total
      J9:M13  copy of it - the rows the report reads: 10 Female, 11 Male,
              12 Corporate, 13 Grand Total
      J15:N20 Net MT Disbursement during the Qtr: GL 1846 balance at the
              previous and current quarter end (N15/N16), N17 = N16 - N15,
              L15 = N17/1000, split by gender in proportion to J9's Sum of Total
    Clients with no gender are left out of both tables (they are listed in
    Missing_Records.xlsx instead).
    """
    ws._pivots = []   # the old template's PivotTable sat on J3:L7
    c_rng, g_rng = f"$C${first}:$C${last}", f"$G${first}:$G${last}"

    rows = [('Female', f'=COUNTIF({c_rng},"Female")', f'=SUMIF({c_rng},"Female",{g_rng})'),
            ('Male', f'=COUNTIF({c_rng},"Male")', f'=SUMIF({c_rng},"Male",{g_rng})'),
            ('Corporate', f'=COUNTIF({c_rng},"Corporate")', f'=SUMIF({c_rng},"Corporate",{g_rng})')]
    grand = 4 + len(rows)

    for c, title in enumerate(['Row Labels', 'Count of Client Name', 'Sum of Total'], start=10):
        cell = ws.cell(3, c, title)
        cell.font, cell.fill = Font(bold=True), SUMMARY_FILL
    for i, (label, count, total) in enumerate(rows):
        ws.cell(4 + i, 10, label)
        ws.cell(4 + i, 11, count)
        ws.cell(4 + i, 12, total).number_format = '#,##0'
    ws.cell(grand, 10, 'Grand Total')
    ws.cell(grand, 11, f'=SUM(K4:K{grand - 1})')
    ws.cell(grand, 12, f'=SUM(L4:L{grand - 1})').number_format = '#,##0'
    for c in range(10, 13):
        cell = ws.cell(grand, c)
        cell.font, cell.fill, cell.border = Font(bold=True), SUMMARY_FILL, Border(top=THIN)

    # Copy of the summary, in the fixed rows the report reads
    for c, title in enumerate(['Row Labels', 'Count of Client Name', 'Sum of Total'], start=10):
        ws.cell(9, c, title)
    ws.cell(9, 13, 1000)
    for r, label, k, l in [(10, 'Female', '=K4', '=L4'), (11, 'Male', '=K5', '=L5'),
                           (12, 'Corporate', '=K6', '=L6'),
                           (13, 'Grand Total', '=SUM(K10:K12)', '=SUM(L10:L12)')]:
        ws.cell(r, 10, label)
        ws.cell(r, 11, k)
        ws.cell(r, 12, l).number_format = '#,##0'

    # Net MT Disbursement during the Qtr (GL 1846)
    gl = gl or {}
    ws.cell(15, 10, 'Net MT Disbursement during the Qtr').font = Font(bold=True)
    ws.cell(15, 12, '=+N17/1000').number_format = '#,##0'
    ws.cell(15, 13, gl.get('prev_label'))
    ws.cell(15, 14, gl.get('prev_value')).number_format = '#,##0.00'
    ws.cell(16, 13, gl.get('cur_label'))
    ws.cell(16, 14, gl.get('cur_value')).number_format = '#,##0.00'
    ws.cell(16, 14).border = Border(bottom=THIN)
    ws.cell(17, 14, '=+N16-N15').number_format = '#,##0'
    ws.cell(17, 14).border = Border(bottom=DOUBLE)
    for r, label, formula in [(17, 'Male', '=L11/L13*L15'), (18, 'Female', '=L10/L13*L15'),
                              (19, 'Other', '=L12/L13*L15'), (20, None, '=SUM(L17:L19)')]:
        ws.cell(r, 10, label).font = Font(bold=True)
        ws.cell(r, 12, formula).number_format = '#,##0'
        for c in range(10, 13):
            ws.cell(r, c).fill = SPLIT_FILL
    ws.cell(19, 12).border = Border(bottom=THIN)
    ws.cell(20, 12).font = Font(bold=True)
    ws.cell(20, 12).border = Border(bottom=DOUBLE)


def _write_missing_records(folder, missing_df):
    """Write clients without a gender to the 'MT' sheet of Missing_Records.xlsx
    (other sheets, e.g. the LTV bot's, are kept). Removes a stale 'MT' sheet
    when nothing is missing."""
    path = Path(folder) / "Missing_Records.xlsx"
    if not path.exists() and missing_df.empty:
        return
    wb = load_workbook(path) if path.exists() else None
    if wb is not None and MISSING_RECORDS_SHEET in wb.sheetnames:
        del wb[MISSING_RECORDS_SHEET]
    if missing_df.empty:
        if wb is not None and wb.sheetnames:
            wb.save(path)
        return
    if wb is None:
        wb = Workbook()
        wb.active.title = MISSING_RECORDS_SHEET
        ws = wb.active
    else:
        ws = wb.create_sheet(MISSING_RECORDS_SHEET)
    ws.append(list(missing_df.columns))
    for row in missing_df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    wb.save(path)
    logger.info(f"{len(missing_df)} MT client(s) without a gender written to {path} ('{MISSING_RECORDS_SHEET}' sheet)")


def consolidate_mt_reports(date, working_dir=None):
    """
    Build the quarter's MT workbook, 'MT - <Mon YYYY>.xlsx' (named after the report month).

    For a 30-Sep-2025 run:
      - Inputs (Input/MT of the last completed outputs folder, where the UI puts the uploads):
          31-07-2025 / 31-08-2025 / 30-09-2025  Client Available Limits Report.xlsx
          the previous run's MT workbook (see _previous_mt_file), e.g. MT - Jul 2025.xlsx
      - Output sheets:
          Jun            - kept from the previous workbook (opening balances)
          Jul, Aug, Sep  - copies of the three reports
          Dis            - MT No., Client Name, Gender, Opening Balance (Jun),
                           Actual Lending 31-07-2025 / 31-08-2025 / 30-09-2025
          Disbursement   - monthly movement per client (Jul - Opening, Aug - Jul,
                           Sep - Aug), Total row and gender summary
      - Saved to <working_dir>/Input/MT (so it is archived with this run and
        becomes the next quarter's previous workbook) and to the source Input/MT.

    Gender here comes only from the previous workbook's Dis sheet; the rest is
    filled in by finalize_mt_workbook() at the end of the report run.

    Returns the path of the saved workbook (the working copy when working_dir is given).
    """
    months = _quarter_months(date)
    opening_month = _month_before(months[0])
    month_labels = [m.strftime('%b') for m in months]
    logger.info(f"MT quarter: {', '.join(m.strftime('%d-%m-%Y') for m in months)}; "
                f"opening balance month: {opening_month.strftime('%b %Y')}")

    # 1. The three monthly reports (read only - the source files are not modified)
    report_paths = [INPUT_FOLDER / f"{m.strftime('%d-%m-%Y')}  Client Available Limits Report.xlsx" for m in months]
    missing = [p.name for p in report_paths if not p.exists()]
    if missing:
        raise FileNotFoundError(f"Client Available Limits report(s) not found in {INPUT_FOLDER}: {', '.join(missing)}")

    report_wbs = [load_workbook(p) for p in report_paths]
    monthly = [read_lending_sheet(wb.worksheets[0], p.name) for wb, p in zip(report_wbs, report_paths)]
    for m, records in zip(months, monthly):
        logger.info(f"  {m.strftime('%d-%m-%Y')}: {len(records)} MT clients")

    # 2. Previous workbook: opening balances (month-before-quarter sheet) and genders (Dis)
    prev_path = _previous_mt_file(months)
    logger.info(f"Previous MT workbook: {prev_path}")
    prev_values = load_workbook(prev_path, data_only=True)
    opening_sheet = opening_month.strftime('%b')
    if opening_sheet in prev_values.sheetnames:
        opening = read_lending_sheet(prev_values[opening_sheet], f"{prev_path.name} [{opening_sheet}]")
        logger.info(f"Opening balances from sheet '{opening_sheet}': {len(opening)} MT clients")
    elif 'Dis' in prev_values.sheetnames:
        logger.warning(f"{prev_path.name} has no '{opening_sheet}' sheet - opening balances taken from "
                       f"the last column of its Dis sheet")
        opening = {k: ('', v) for k, v in _previous_dis_closing(prev_values['Dis']).items()}
    else:
        raise ValueError(f"{prev_path.name} has neither a '{opening_sheet}' sheet nor a Dis sheet")
    genders = _previous_genders(prev_values['Dis']) if 'Dis' in prev_values.sheetnames else {}

    # 3. Dis: every MT No. from the 4 sheets once; latest available name
    all_mt = set(opening)
    for records in monthly:
        all_mt |= set(records)
    rows = []
    for mt_no in sorted(all_mt, key=_mt_sort_key):
        name = next((src[mt_no][0] for src in reversed(monthly + [opening]) if mt_no in src and src[mt_no][0]), '')
        rows.append([mt_no, name, genders.get(mt_no), opening.get(mt_no, ('', 0.0))[1]]
                    + [records.get(mt_no, ('', 0.0))[1] for records in monthly])
    lending_cols = [f"Actual Lending {m.strftime('%d-%m-%Y')}" for m in months]
    dis_df = pd.DataFrame(rows, columns=['MT No.', 'Client Name', 'Gender', 'Opening Balance'] + lending_cols)
    logger.info(f"Dis: {len(dis_df)} unique MT clients")

    # 4. Build the workbook from the previous one: keep only the opening month
    #    sheet and the Disbursement sheet (as the template), then add the rest.
    wb = load_workbook(prev_path)
    for name in list(wb.sheetnames):
        if name not in (opening_sheet, 'Disbursement'):
            del wb[name]
    for m, report_wb in zip(months, report_wbs):
        label = m.strftime('%b')
        if label in wb.sheetnames:
            del wb[label]
        _copy_report_to_sheet(report_wb.worksheets[0], wb.create_sheet(label))
    _write_dis_sheet(wb, dis_df)
    _write_disbursement_sheet(wb, dis_df, month_labels)

    # 5. Save: this run's working Input/MT (archived with the run) and the source Input/MT.
    out_name = f"MT - {months[-1].strftime('%b %Y')}.xlsx"
    if working_dir:
        working_mt = Path(working_dir) / "Input" / "MT"
        working_mt.mkdir(parents=True, exist_ok=True)
        main_path = working_mt / out_name
    else:
        main_path = INPUT_FOLDER / out_name
    _save_mt_workbook(wb, main_path, [opening_sheet] + month_labels)

    # 6. Clients with no gender -> Missing_Records.xlsx in this run's folder
    #    (finalize_mt_workbook() fills most of them in at the end of the run)
    missing_df = dis_df[dis_df['Gender'].isna()]
    _write_missing_records(working_dir or OUTPUT_FOLDER, missing_df)

    return main_path


def _save_mt_workbook(wb, main_path, month_sheets):
    """Order the sheets, save to main_path, let Excel recalculate it (openpyxl
    writes formulas without results - the values are needed when the
    Disbursement sheet is copied into the report as values) and copy it to the
    source Input/MT folder too."""
    order = ['Disbursement'] + month_sheets + ['Dis']
    wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else len(order))
    wb.active = 0
    wb.calculation.calcMode = 'auto'
    wb.calculation.fullCalcOnLoad = True
    main_path = Path(main_path)
    wb.save(main_path)
    try:
        _recalculate_with_excel(main_path)
    except Exception as e:
        logger.warning(f"Excel recalculation of {main_path.name} failed ({e}); formula cells have no stored values")
    logger.info(f"Saved {main_path}")
    source_copy = INPUT_FOLDER / main_path.name
    if main_path.resolve() != source_copy.resolve():
        shutil.copy2(main_path, source_copy)
        logger.info(f"Saved {source_copy}")


MT_CLIENT_CODES_DIR = parent_directory / "Master Inputs" / "MT Client Codes"


def _normalize_code(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def load_mt_client_codes(folder=None):
    """{MT No.: ICAM code} from the latest uploaded 'MT Client Codes' file
    (columns 'ICAM Code' and 'MT', any order, header in the first rows)."""
    folder = Path(folder or MT_CLIENT_CODES_DIR)
    files = [f for f in folder.glob("*.xlsx") if not f.name.startswith("~$")] if folder.exists() else []
    if not files:
        logger.warning(f"No MT Client Codes file in {folder} - ICAM codes unknown")
        return {}
    path = max(files, key=os.path.getmtime)
    ws = load_workbook(path, data_only=True).worksheets[0]
    icam_hdr, mt_hdr = _find_header(ws, 'ICAM Code'), _find_header(ws, 'MT')
    if not (icam_hdr and mt_hdr):
        logger.warning(f"{path.name}: 'ICAM Code' / 'MT' headers not found - ICAM codes unknown")
        return {}
    codes = {}
    for r in range(max(icam_hdr[0], mt_hdr[0]) + 1, ws.max_row + 1):
        mt_no, icam = ws.cell(r, mt_hdr[1]).value, _normalize_code(ws.cell(r, icam_hdr[1]).value)
        if mt_no and icam and MT_NO_RE.match(str(mt_no).strip()):
            codes[str(mt_no).strip().upper()] = icam
    logger.info(f"MT Client Codes: {len(codes)} MT No. -> ICAM code from {path.name}")
    return codes


def _report_genders(report_wb):
    """{client code: Male/Female} from the report's Disbursement sheet, i.e. the
    genders this run already resolved (title formula + disbursement_gap_fill_bot
    ERP results). A formula cell is evaluated the same way the formula does it."""
    from bots.disbursement_gap_fill_bot import gender_from_name
    genders = {}
    if 'Disbursement' not in report_wb.sheetnames:
        return genders
    ws = report_wb['Disbursement']
    for r in range(5, ws.max_row + 1):
        code = _normalize_code(ws.cell(r, 1).value)
        if not code:
            continue
        value = ws.cell(r, 6).value
        if isinstance(value, str) and value.startswith('='):
            value = gender_from_name(ws.cell(r, 3).value)
        if value in ('Male', 'Female'):
            genders.setdefault(code, value)
    return genders


def finalize_mt_workbook(mt_path, report_wb, report_date, records_folder):
    """End-of-run step for the MT workbook built by consolidate_mt_reports(),
    then copy its Disbursement sheet into the report's 'MT-Disbursement' (values).

    1. Genders still blank in Dis (none in the previous MT workbook), via the
       MT No. -> ICAM code upload:
         ICAM code '2...'         -> Corporate
         in the report's Disbursement sheet (this run's own gender results)
         otherwise                -> ERP Customer Inquiry
       Still unknown: left blank (left out of the gender summary) and
       listed in Missing_Records.xlsx.
    2. GL 1846 balance at the previous and current quarter end from the ERP
       (opening balance on the 1st of the following month) -> N15 / N16.
    3. Rebuild Dis + Disbursement, save (Excel recalculates), copy to the report.
    report_wb is an open openpyxl workbook; the caller saves it.
    """
    if str(parent_directory) not in sys.path:
        sys.path.insert(0, str(parent_directory))   # for the bots package
    mt_path = Path(mt_path)
    months = _quarter_months(report_date)
    month_labels = [m.strftime('%b') for m in months]
    opening_sheet = _month_before(months[0]).strftime('%b')

    wb = load_workbook(mt_path)
    dis_df = pd.read_excel(mt_path, sheet_name='Dis')
    dis_df['Gender'] = dis_df['Gender'].where(dis_df['Gender'].notna(), None)

    # 1. Genders
    icam_codes = load_mt_client_codes()
    memory = _report_genders(report_wb)
    blank = dis_df['Gender'].isna()
    logger.info(f"MT genders: {int(blank.sum())} blank after the previous MT workbook; "
                f"{len(memory)} client genders known from this run's Disbursement sheet")
    to_search = {}
    for i in dis_df.index[blank]:
        icam = icam_codes.get(str(dis_df.at[i, 'MT No.']).strip().upper())
        if not icam:
            continue
        if icam.startswith('2'):
            dis_df.at[i, 'Gender'] = 'Corporate'
        elif icam in memory:
            dis_df.at[i, 'Gender'] = memory[icam]
        elif icam.startswith('1'):
            to_search[i] = icam
    if to_search:
        try:
            from bots.mt_erp_lookups import fetch_genders
            found = fetch_genders(to_search.values())
            for i, icam in to_search.items():
                if found.get(icam):
                    dis_df.at[i, 'Gender'] = found[icam]
        except Exception as e:
            logger.warning(f"ERP gender lookup failed ({e}); those clients stay blank")
    still_blank = dis_df['Gender'].isna()
    logger.info(f"MT genders: {int((blank & ~still_blank).sum())} filled, {int(still_blank.sum())} still blank")

    # 2. GL 1846 balances
    prev_end, cur_end = _month_before(months[0]), months[-1]
    prev_date = (prev_end + timedelta(days=1)).strftime('%d/%m/%Y')
    cur_date = (cur_end + timedelta(days=1)).strftime('%d/%m/%Y')
    gl = {'prev_label': prev_end.strftime('%b'), 'cur_label': cur_end.strftime('%b')}
    try:
        from bots.mt_erp_lookups import fetch_gl_opening_balances
        balances = fetch_gl_opening_balances([prev_date, cur_date])
        gl['prev_value'], gl['cur_value'] = balances.get(prev_date), balances.get(cur_date)
    except Exception as e:
        logger.warning(f"GL 1846 balance lookup failed ({e}); N15/N16 left blank")

    # 3. Rebuild, save, copy into the report
    _write_dis_sheet(wb, dis_df)
    _write_disbursement_sheet(wb, dis_df, month_labels, gl)
    _save_mt_workbook(wb, mt_path, [opening_sheet] + month_labels)

    missing = dis_df[still_blank][['MT No.', 'Client Name']].copy()
    missing['ICAM Code'] = [icam_codes.get(str(m).upper()) for m in missing['MT No.']]
    missing['Reason'] = ['ICAM code not in MT Client Codes upload' if not c else 'Gender not found (ERP)'
                         for c in missing['ICAM Code']]
    _write_missing_records(records_folder, missing)

    if 'MT-Disbursement' in report_wb.sheetnames:
        copy_sheet_as_values(mt_path, 'Disbursement', report_wb['MT-Disbursement'])
        logger.info(f"Copied {mt_path.name} Disbursement -> report 'MT-Disbursement'")
    else:
        logger.warning("'MT-Disbursement' sheet not found in the report - not copied")


def _recalculate_with_excel(path):
    """Open the workbook in Excel, recalculate and save, so every formula cell
    also stores its result."""
    import xlwings as xw
    app = xw.App(visible=False, add_book=False)
    app.display_alerts = False
    try:
        book = app.books.open(str(path))
        app.calculate()
        book.save()
        book.close()
    finally:
        app.quit()


def copy_sheet_as_values(src_path, sheet_name, dst_ws):
    """Replace dst_ws with a copy of `sheet_name` from the workbook at src_path:
    cell values (formula results, not formulas), fonts, fills, borders,
    alignment, number formats, merged cells, column widths and row heights."""
    styles_ws = load_workbook(src_path)[sheet_name]
    values_ws = load_workbook(src_path, data_only=True)[sheet_name]

    for merged in list(dst_ws.merged_cells.ranges):
        dst_ws.unmerge_cells(str(merged))
    if dst_ws.max_row:
        dst_ws.delete_rows(1, dst_ws.max_row)   # drops old values and styles; other sheets' references stay as typed

    for row in styles_ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            value = values_ws.cell(cell.row, cell.column).value
            if isinstance(value, str) and value.startswith('='):
                value = None
            target = dst_ws.cell(row=cell.row, column=cell.column)
            target.value = value
            if cell.has_style:
                target.font = copy(cell.font)
                target.fill = copy(cell.fill)
                target.border = copy(cell.border)
                target.alignment = copy(cell.alignment)
                target.number_format = cell.number_format
                target.protection = copy(cell.protection)

    for merged in styles_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merged))
    for key, dim in styles_ws.column_dimensions.items():
        dst_ws.column_dimensions[key].width = dim.width
        dst_ws.column_dimensions[key].hidden = dim.hidden
    for key, dim in styles_ws.row_dimensions.items():
        if dim.height is not None:
            dst_ws.row_dimensions[key].height = dim.height
    dst_ws.freeze_panes = styles_ws.freeze_panes
    # Only the sheet's contents are copied, not its tab colour; clear any tab
    # colour an earlier version of this copy left behind (the report's
    # template sheet has none).
    dst_ws.sheet_properties.tabColor = None


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format='%(asctime)s [%(levelname)s] %(message)s', datefmt='%H:%M:%S')
    # Execute the consolidation using the same --report-date argument process_summary.py accepts
    print(f"Using report date: {target_report_date.strftime('%Y-%m-%d')}")
    result_path = consolidate_mt_reports(target_report_date)
    print(f"\nMT workbook saved to: {result_path}")
