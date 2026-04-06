#!/usr/bin/env python3
r"""
Monthly Liquidity Report Automation Script
Central Bank Sri Lanka - NBD_MF_04_LA

Consolidates 5 AFL Liquidity files into a single monthly NBD-MF-04-LA report.

USAGE:
    python NBD_MF_04_LA.py MM/DD/YYYY

WORKFLOW:
1. Find working folder: working/NBD_MF_04_LA/{single-date-subfolder}/
2. Find template: NBD-MF-04-LA*.xlsm in working folder
3. Find 5 AFL files from working folder:
   - AFL Liquidity - Week_1.xlsx  -> Week 1 sheet (paste C3:C22 values to H1:H20)
   - AFL Liquidity - Week_2.xlsx  -> Week 2 sheet (paste C3:C22 values to H1:H20)
   - AFL Liquidity - Week_3.xlsx  -> Week 3 sheet (paste C3:C22 values to H1:H20)
   - AFL Liquidity - Week_4.xlsx  -> Week 4 sheet (paste C3:C22 values to H1:H20)
   - AFL Liquidity - Mon-YYYY.xlsx -> Week 5 sheet (paste C3:C22 values to J1:J20)
4. From each AFL file, read "NBD-WF-15-LA" sheet C3:C22 (values only, no formulas)
5. Paste values into the corresponding Week sheet at the target range
6. Fill NBD-MF-04-LA sheet columns C-H with data from Week sheets
7. Save monthly report in working folder

SOURCE RANGE:  NBD-WF-15-LA sheet  C3:C22  (20 rows, values only)
TARGET RANGES:
  Week 1-4 sheets -> H1:H20
  Week 5 sheet    -> J1:J20
"""

import os
import sys
import logging
from datetime import datetime
from pathlib import Path

import openpyxl

# Setup logging
def setup_logger(script_dir):
    logs_dir = os.path.join(script_dir, "logs")
    os.makedirs(logs_dir, exist_ok=True)

    log_file = os.path.join(
        logs_dir,
        f"NBD_MF_04_LA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
    )

    root_logger = logging.getLogger()
    root_logger.setLevel(logging.INFO)
    root_logger.handlers.clear()

    file_handler = logging.FileHandler(log_file, encoding='utf-8')
    file_handler.setLevel(logging.INFO)
    file_handler.setFormatter(logging.Formatter("%(asctime)s - %(levelname)s - %(message)s"))

    console_handler = logging.StreamHandler(sys.stdout)
    console_handler.setLevel(logging.INFO)
    console_handler.setFormatter(logging.Formatter("%(asctime)s - %(levelname)s - %(message)s"))

    root_logger.addHandler(file_handler)
    root_logger.addHandler(console_handler)

    logging.info("Logger initialized")
    print(f"Logger initialized. Log file: {log_file}")
    return log_file

logger = logging.getLogger(__name__)


def find_single_subdirectory(parent_dir):
    """Find the single subdirectory within a parent directory."""
    if not parent_dir.exists():
        logger.error(f"Parent directory does not exist: {parent_dir}")
        return None

    subdirs = [d for d in parent_dir.iterdir() if d.is_dir()]

    if len(subdirs) == 0:
        logger.error(f"No subdirectories found in {parent_dir}")
        return None
    elif len(subdirs) > 1:
        logger.warning(f"Multiple subdirectories found in {parent_dir}, using first one: {subdirs[0].name}")
        return subdirs[0]
    else:
        logger.info(f"Found single subdirectory: {subdirs[0].name}")
        return subdirs[0]


def find_afl_files(working_folder):
    """
    Find 5 AFL Liquidity files in the working folder.
    Returns list of 5 paths in order: Week 1, Week 2, Week 3, Week 4, Month-End (Week 5).
    """
    afl_files = []

    logger.info("Looking for AFL files:")

    # Find Week 1-4 files in exact order
    for week in range(1, 5):
        week_file = working_folder / f"AFL Liquidity - Week_{week}.xlsx"
        if week_file.exists():
            afl_files.append(week_file)
            logger.info(f"  Week {week}: {week_file.name} ✓")
        else:
            logger.warning(f"  Week {week}: NOT FOUND — expected '{week_file.name}'")

    # Find month-end file: AFL Liquidity - Mon-YYYY.xlsx (no "Week_" in name)
    month_end_files = [
        f for f in working_folder.glob("AFL Liquidity - *-*.xlsx")
        if "Week_" not in f.name
    ]

    if month_end_files:
        month_end_file = max(month_end_files, key=lambda f: f.stat().st_mtime)
        afl_files.append(month_end_file)
        logger.info(f"  Week 5 (Month End): {month_end_file.name} ✓")
    else:
        logger.warning("  Week 5 (Month End): NOT FOUND — expected format: AFL Liquidity - Mon-YYYY.xlsx")

    logger.info(f"Total AFL files found: {len(afl_files)} / 5")
    return afl_files


def read_afl_source_values(afl_file_path):
    """
    Read calculated values from the "NBD-WF-15-LA" sheet, range C3:C22 (20 rows).
    Uses Excel COM automation to force formula recalculation before reading,
    ensuring actual values are returned rather than None for formula cells.
    Falls back to openpyxl data_only if COM is unavailable.
    Returns a list of 20 values.
    """
    file_path_str = str(afl_file_path.resolve())
    logger.info(f"  Opening via Excel COM: {afl_file_path.name}")

    excel = None
    wb_com = None
    try:
        import win32com.client
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        # Open workbook, trigger full recalculation, then read
        wb_com = excel.Workbooks.Open(
            file_path_str,
            UpdateLinks=False,
            ReadOnly=True
        )
        excel.CalculateFull()

        # Find the sheet
        ws_com = None
        for sh in wb_com.Sheets:
            if sh.Name == "NBD-WF-15-LA":
                ws_com = sh
                break

        if ws_com is None:
            available = [sh.Name for sh in wb_com.Sheets]
            logger.error(f"  Sheet 'NBD-WF-15-LA' not found. Available: {available}")
            return None

        # Read C3:C22 (column 3, rows 3-22)
        values = []
        for row in range(3, 23):
            cell_val = ws_com.Cells(row, 3).Value
            values.append(cell_val)

        logger.info(f"  Read {len(values)} values from {afl_file_path.name} [NBD-WF-15-LA] C3:C22")
        for i, v in enumerate(values, start=3):
            logger.info(f"    C{i} = {v}")

        return values

    except Exception as e:
        logger.warning(f"  Excel COM failed ({e}), falling back to openpyxl data_only...")
        # Fallback: openpyxl with data_only (cached values only — may return None for formulas)
        wb = None
        try:
            wb = openpyxl.load_workbook(afl_file_path, data_only=True, read_only=True)
            if "NBD-WF-15-LA" not in wb.sheetnames:
                logger.error(f"  Sheet 'NBD-WF-15-LA' not found in {afl_file_path.name}")
                return None
            ws = wb["NBD-WF-15-LA"]
            values = [ws.cell(row=r, column=3).value for r in range(3, 23)]
            logger.info(f"  Read {len(values)} values (openpyxl fallback) from {afl_file_path.name}")
            for i, v in enumerate(values, start=3):
                logger.info(f"    C{i} = {v}")
            return values
        except Exception as e2:
            logger.error(f"  openpyxl fallback also failed: {e2}")
            return None
        finally:
            if wb:
                wb.close()
    finally:
        if wb_com is not None:
            try:
                wb_com.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass


def paste_values_to_week_sheet(monthly_wb, week_sheet_name, values, target_col):
    """
    Paste 20 values into the week sheet at the target column, rows 1-20.

    Args:
        monthly_wb: the monthly workbook
        week_sheet_name: e.g. "Week 1"
        values: list of 20 values
        target_col: column letter, e.g. "H" for weeks 1-4, "J" for week 5
    """
    if week_sheet_name not in monthly_wb.sheetnames:
        logger.error(f"Sheet '{week_sheet_name}' not found in monthly report")
        logger.error(f"  Available sheets: {monthly_wb.sheetnames}")
        return False

    ws = monthly_wb[week_sheet_name]

    # Paste values into target_col rows 1-20
    for i, value in enumerate(values):
        row = i + 1  # row 1 to 20
        ws[f"{target_col}{row}"].value = value

    logger.info(f"  Pasted {len(values)} values into {week_sheet_name}!{target_col}1:{target_col}20")
    return True


def fill_nbd_mf_04_la_sheet(monthly_wb):
    """
    Fill NBD-MF-04-LA sheet columns C-H with values from Week 1-5 sheets.

    Mapping for columns C-F (use H column from Week sheets):
    - C5:C15 <- Week 1 H3:H13, C19 <- Week 1 H14, C22:C27 <- Week 1 H15:H20
    - Same pattern for D (Week 2), E (Week 3), F (Week 4)

    Mapping for columns G-H (both use Week 5 J column):
    - G5:G15 <- Week 5 J3:J13, G19 <- Week 5 J14, G22:G27 <- Week 5 J15:J20
    - H5:H15 <- Week 5 J3:J13, H19 <- Week 5 J14, H22:H27 <- Week 5 J15:J20
    """
    try:
        logger.info("\n--- Filling NBD-MF-04-LA Sheet ---")

        if "NBD-MF-04-LA" not in monthly_wb.sheetnames:
            logger.error("Sheet 'NBD-MF-04-LA' not found in monthly report")
            logger.info(f"Available sheets: {monthly_wb.sheetnames}")
            return False

        nbd_sheet = monthly_wb["NBD-MF-04-LA"]

        week_to_column = {
            1: 'C',
            2: 'D',
            3: 'E',
            4: 'F',
        }

        # Process weeks 1-4: read H column from each Week sheet -> paste to NBD-MF-04-LA
        for week_num in range(1, 5):
            week_sheet_name = f"Week {week_num}"

            if week_sheet_name not in monthly_wb.sheetnames:
                logger.warning(f"Sheet '{week_sheet_name}' not found, skipping...")
                continue

            week_sheet = monthly_wb[week_sheet_name]
            dest_col = week_to_column[week_num]

            logger.info(f"  Week {week_num}: {week_sheet_name}!H -> NBD-MF-04-LA!{dest_col}")

            # H3:H13 -> dest_col 5:15 (11 cells)
            for i in range(11):
                nbd_sheet[f"{dest_col}{5 + i}"].value = week_sheet[f"H{3 + i}"].value

            # H14 -> dest_col 19 (1 cell)
            nbd_sheet[f"{dest_col}19"].value = week_sheet["H14"].value

            # H15:H20 -> dest_col 22:27 (6 cells)
            for i in range(6):
                nbd_sheet[f"{dest_col}{22 + i}"].value = week_sheet[f"H{15 + i}"].value

            logger.info(f"    Done")

        # Process Week 5: read J column -> paste to both G and H columns in NBD-MF-04-LA
        week5_sheet_name = "Week 5"
        if week5_sheet_name in monthly_wb.sheetnames:
            week5_sheet = monthly_wb[week5_sheet_name]
            logger.info(f"  Week 5: {week5_sheet_name}!J -> NBD-MF-04-LA!G and H")

            for dest_col in ['G', 'H']:
                # J3:J13 -> dest_col 5:15 (11 cells)
                for i in range(11):
                    nbd_sheet[f"{dest_col}{5 + i}"].value = week5_sheet[f"J{3 + i}"].value

                # J14 -> dest_col 19 (1 cell)
                nbd_sheet[f"{dest_col}19"].value = week5_sheet["J14"].value

                # J15:J20 -> dest_col 22:27 (6 cells)
                for i in range(6):
                    nbd_sheet[f"{dest_col}{22 + i}"].value = week5_sheet[f"J{15 + i}"].value

            logger.info(f"    Done")
        else:
            logger.warning(f"Sheet 'Week 5' not found — columns G and H not filled")

        logger.info("NBD-MF-04-LA sheet filled successfully")
        return True

    except Exception as e:
        logger.error(f"Error filling NBD-MF-04-LA sheet: {e}")
        import traceback
        logger.error(traceback.format_exc())
        return False


def main():
    """Main execution flow for monthly liquidity report generation."""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    setup_logger(script_dir)

    # Parse command line arguments
    if len(sys.argv) < 2:
        print("Usage: python NBD_MF_04_LA.py MM/DD/YYYY")
        sys.exit(1)

    date_formats = ["%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"]
    input_date = None
    for fmt in date_formats:
        try:
            input_date = datetime.strptime(sys.argv[1], fmt)
            break
        except ValueError:
            continue

    if input_date is None:
        print(f"Error: Could not parse date '{sys.argv[1]}'")
        sys.exit(1)

    logger.info("=" * 60)
    logger.info("NBD-MF-04-LA Monthly Liquidity Report Generation")
    logger.info(f"Report date: {input_date.strftime('%m/%d/%Y')}")
    logger.info("=" * 60)

    # Step 1: Find working folder
    logger.info("\nStep 1: Finding working folder...")
    script_dir = Path(__file__).parent
    base_working = script_dir.parent / "working" / "NBD_MF_04_LA"
    working_folder = find_single_subdirectory(base_working)

    if working_folder is None:
        logger.error(f"Could not find working folder under {base_working}")
        sys.exit(1)

    logger.info(f"Working folder: {working_folder}")

    # Step 2: Find template file
    logger.info("\nStep 2: Finding monthly report template...")
    template_files = list(working_folder.glob("NBD-MF-04-LA*.xlsm"))

    if not template_files:
        logger.error(f"No NBD-MF-04-LA*.xlsm template found in {working_folder}")
        sys.exit(1)

    template_path = max(template_files, key=lambda f: f.stat().st_mtime)
    logger.info(f"Found template: {template_path.name}")

    # Step 3: Find AFL files
    logger.info("\nStep 3: Finding AFL Liquidity files from working folder...")
    afl_files = find_afl_files(working_folder)

    if len(afl_files) < 5:
        logger.error(f"Need exactly 5 AFL files, only found {len(afl_files)}")
        logger.error("Expected:")
        logger.error("  AFL Liquidity - Week_1.xlsx")
        logger.error("  AFL Liquidity - Week_2.xlsx")
        logger.error("  AFL Liquidity - Week_3.xlsx")
        logger.error("  AFL Liquidity - Week_4.xlsx")
        logger.error("  AFL Liquidity - Mon-YYYY.xlsx (month-end)")
        sys.exit(1)

    # Step 4-6: Open monthly file via Excel COM, write values, save, rename
    logger.info("\nStep 4: Opening monthly report via Excel COM...")

    report_month = input_date.strftime("%b")  # e.g., "Feb"
    report_year  = input_date.strftime("%y")  # e.g., "26"
    new_filename = f"NBD-MF-04-LA Liquid Assets - {report_month} {report_year}.xlsm"

    # Week configs: (week_num, sheet_name, target_col_index)
    # H = column 8, J = column 10
    week_configs = [
        (1, "Week 1", 8),   # H
        (2, "Week 2", 8),   # H
        (3, "Week 3", 8),   # H
        (4, "Week 4", 8),   # H
        (5, "Week 5", 10),  # J
    ]

    # Read all AFL values first (closes each COM session before opening monthly)
    logger.info("\nReading AFL source values...")
    all_values = []
    for idx, (week_num, sheet_name, target_col) in enumerate(week_configs):
        afl_file = afl_files[idx]
        logger.info(f"\n--- Week {week_num}: {afl_file.name} ---")
        values = read_afl_source_values(afl_file)
        if values is None:
            logger.error(f"  Failed to read values from {afl_file.name}")
            sys.exit(1)
        all_values.append(values)

    # Now open the monthly file via COM and write all values
    import win32com.client
    excel = None
    monthly_wb_com = None
    try:
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        monthly_path_str = str(template_path.resolve())
        logger.info(f"\nStep 5: Opening monthly file via COM: {template_path.name}")
        monthly_wb_com = excel.Workbooks.Open(monthly_path_str, UpdateLinks=False)

        # Paste values into each Week sheet
        logger.info("\nStep 5a: Pasting values into Week sheets...")
        for idx, (week_num, sheet_name, target_col) in enumerate(week_configs):
            values = all_values[idx]

            # Find sheet by name
            ws = None
            for sh in monthly_wb_com.Sheets:
                if sh.Name == sheet_name:
                    ws = sh
                    break

            if ws is None:
                available = [sh.Name for sh in monthly_wb_com.Sheets]
                logger.error(f"  Sheet '{sheet_name}' not found. Available: {available}")
                continue

            # Paste 20 values into target column, rows 1-20
            for i, value in enumerate(values):
                row = i + 1
                ws.Cells(row, target_col).Value = value

            col_letter = "H" if target_col == 8 else "J"
            logger.info(f"  {sheet_name}!{col_letter}1:{col_letter}20 written ({week_num} of 5)")

        # Fill NBD-MF-04-LA sheet from Week sheets using COM
        logger.info("\nStep 5b: Filling NBD-MF-04-LA sheet...")
        nbd_sheet = None
        for sh in monthly_wb_com.Sheets:
            if sh.Name == "NBD-MF-04-LA":
                nbd_sheet = sh
                break

        if nbd_sheet is None:
            logger.error("Sheet 'NBD-MF-04-LA' not found in monthly report")
        else:
            week_to_col = {1: 3, 2: 4, 3: 5, 4: 6}  # C=3, D=4, E=5, F=6

            for week_num in range(1, 5):
                week_sh = None
                for sh in monthly_wb_com.Sheets:
                    if sh.Name == f"Week {week_num}":
                        week_sh = sh
                        break
                if week_sh is None:
                    logger.warning(f"  Week {week_num} sheet not found, skipping")
                    continue

                dest_col = week_to_col[week_num]
                # H3:H13 -> dest_col rows 5-15
                for i in range(11):
                    nbd_sheet.Cells(5 + i, dest_col).Value = week_sh.Cells(3 + i, 8).Value
                # H14 -> dest_col row 19
                nbd_sheet.Cells(19, dest_col).Value = week_sh.Cells(14, 8).Value
                # H15:H20 -> dest_col rows 22-27
                for i in range(6):
                    nbd_sheet.Cells(22 + i, dest_col).Value = week_sh.Cells(15 + i, 8).Value

                logger.info(f"  Week {week_num} -> NBD-MF-04-LA col {dest_col} done")

            # Week 5 J column -> columns G(7) and H(8) of NBD-MF-04-LA
            week5_sh = None
            for sh in monthly_wb_com.Sheets:
                if sh.Name == "Week 5":
                    week5_sh = sh
                    break
            if week5_sh is None:
                logger.warning("  Week 5 sheet not found, columns G and H not filled")
            else:
                for dest_col in [7, 8]:  # G=7, H=8
                    for i in range(11):
                        nbd_sheet.Cells(5 + i, dest_col).Value = week5_sh.Cells(3 + i, 10).Value
                    nbd_sheet.Cells(19, dest_col).Value = week5_sh.Cells(14, 10).Value
                    for i in range(6):
                        nbd_sheet.Cells(22 + i, dest_col).Value = week5_sh.Cells(15 + i, 10).Value
                logger.info("  Week 5 -> NBD-MF-04-LA cols G and H done")

        # Save the file in place (overwrite existing)
        logger.info(f"\nStep 6: Saving file (overwrite)...")
        monthly_wb_com.Save()
        logger.info(f"  Saved: {template_path.name}")

        # Rename if the filename differs from the target
        new_path = template_path.parent / new_filename
        if template_path.resolve() != new_path.resolve():
            monthly_wb_com.Close(SaveChanges=False)
            monthly_wb_com = None
            if new_path.exists():
                new_path.unlink()
            template_path.rename(new_path)
            logger.info(f"  Renamed to: {new_filename}")
        else:
            monthly_wb_com.Close(SaveChanges=False)
            monthly_wb_com = None
            logger.info(f"  Filename unchanged: {new_filename}")

    finally:
        if monthly_wb_com is not None:
            try:
                monthly_wb_com.Close(SaveChanges=False)
            except Exception:
                pass
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass

    logger.info("\n" + "=" * 60)
    logger.info("SUCCESS: Monthly Liquidity Report generated successfully")
    logger.info(f"Output: {template_path.parent / new_filename}")
    logger.info("=" * 60)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        logger.error(f"Error in monthly report generation: {e}")
        import traceback
        logger.error(traceback.format_exc())
        sys.exit(1)
