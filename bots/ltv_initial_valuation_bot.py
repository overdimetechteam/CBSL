#!/usr/bin/env python3
"""
Fills #N/A 'Initial Valuation' cells in the report's 'LTV Working - Disbursement'
sheet from the ERP (Recovery -> Information Center -> Contract Inquiry).

Per contract:
  1. Search the Contract No, open Equipment Details, read Valuation Amount.
  2. If that is empty/0: type the contract's Client Code (no Enter), click the
     "?" icon after Contract No, take the client's other Active ('A') contracts
     from the list, and try each one's Equipment Details the same way.
  3. If still empty/0: the row goes to a 'LTV Working - Disbursement' sheet in
     Missing_Records.xlsx (next to the report).

Run it yourself and watch individual records (visible Chrome, every step printed):
    python bots/ltv_initial_valuation_bot.py --contracts WLAL260159320,WLAT260121470
Run against a report file (what process_summary.py does):
    python bots/ltv_initial_valuation_bot.py --report-file "<Quarter Classification ... .xlsx>"
"""
import argparse
import json
import logging
import os
import sys
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from selenium.common.exceptions import TimeoutException, WebDriverException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
# Reuses the ERP login and popup-dismissal helpers already proven in the
# Disbursement gap-fill bot (same credential scheme, same ERP dialogs).
from disbursement_gap_fill_bot import (  # noqa: E402
    _erp_login, _dismiss_stray_alert, _dismiss_error_dialog, _is_connection_error,
)

logger = logging.getLogger("ltv_initial_valuation_bot")

RECOVERY_URL = "https://erp.assetline.lk/Application/Home/RECOVERY"
CONTRACT_INQUIRY_URL = "https://erp.assetline.lk/Recovery/RECContractInquiry"
INQUIRE_API_URL = "https://erp.assetline.lk/Recovery/RECContractInquiry/Inquire"

ID_ADD_BUTTON = "btn_02461_Add"
ID_CONTRACT_INPUT = "strContractNo_RECContractInquiry"
ID_CLIENT_CODE_INPUT = "strClientCode_RECContractInquiry"
ID_CONTRACT_HELP_ICON = "hlp_02461_ContractNo_RECContractInquiry"
ID_EQUIPMENT_BUTTON = "btn_02461_EquipmentDetails_Button_RECContractInquiry"
ID_EQUIPMENT_CLOSE = "btn_02461_Equipment_DetailsBtnClose"
ID_HELP_TABLE_BODY = "thbody02461"
ID_HELP_CLOSE = "btnClose"

# Contract list columns (from the popup's table): Contract No, Client Code, ..., Status
HELP_COL_CONTRACT = 0
HELP_COL_STATUS = 7

LTV_SHEET = "LTV Working - Disbursement"
LTV_START_ROW = 5
COL_CONTRACT, COL_NAME, COL_PRODUCT, COL_EQUIPMENT, COL_GROSS = 1, 2, 3, 4, 5
COL_COLLATERAL, COL_GRANT, COL_INIT_VAL, COL_LTV, COL_WALTV = 6, 7, 8, 9, 10
LTV_HEADERS = ['Contract No', 'Customer Name', 'Product', 'Equipment', 'Gross Exposure',
               'Collateral Type', 'Grant Amount', 'Initial Valuation', 'LTV %', 'WALTV %']

SUMMARY_SHEET = "Summary"
SUMMARY_COL_INITIAL_VALUATION = 51  # column AY - what the sheet's VLOOKUP returns
DISBURSEMENT_SHEET = "Disbursement"
DISBURSEMENT_START_ROW = 5


# ---------------------------------------------------------------- ERP steps

def _pause(seconds):
    time.sleep(seconds)


def navigate_to_contract_inquiry(driver, wait):
    """Same route as bots/IA_Working_Initial_valuation_bot.py: Recovery module ->
    Information Center tile -> Contract Inquiry tile -> Add; direct URL fallback."""
    print("Opening Recovery module -> Information Center -> Contract Inquiry...")
    driver.get(RECOVERY_URL)
    _pause(3)
    try:
        info_center = wait.until(EC.element_to_be_clickable((
            By.XPATH, "//div[contains(@class,'metro-tile-body') and contains(@class,'tile-type-menu')]"
                      "//i[contains(@class,'fa-file')]/..")))
        info_center.click()
        _pause(3)
        inquiry = wait.until(EC.element_to_be_clickable((
            By.XPATH, "//div[contains(@class,'metro-tile-body')]//i[contains(@class,'fa-user')]/..")))
        inquiry.click()
        _pause(3)
    except Exception as e:
        logger.warning(f"Tile navigation failed ({type(e).__name__}); using direct Contract Inquiry URL")
        driver.get(CONTRACT_INQUIRY_URL)
    wait.until(EC.element_to_be_clickable((By.ID, ID_ADD_BUTTON))).click()
    _pause(2)
    print("On Contract Inquiry screen.\n")


def _reset_form(driver):
    """Click Add to blank the inquiry form (best effort)."""
    try:
        btn = driver.find_element(By.ID, ID_ADD_BUTTON)
        driver.execute_script("arguments[0].click();", btn)
        _pause(1.5)
    except Exception:
        pass


def _type_into(driver, wait, element_id, text):
    field = wait.until(EC.element_to_be_clickable((By.ID, element_id)))
    try:
        field.clear()
    except Exception:
        pass
    field.send_keys(text)
    return field


def _search_contract(driver, wait, contract_no):
    """Type the Contract No and press Enter. Returns the Client Code the ERP
    auto-filled (may be '' if the contract wasn't found)."""
    _dismiss_stray_alert(driver, contract_no)
    _dismiss_error_dialog(driver, contract_no)
    field = _type_into(driver, wait, ID_CONTRACT_INPUT, contract_no)
    field.send_keys("\n")
    _pause(3)
    return driver.find_element(By.ID, ID_CLIENT_CODE_INPUT).get_attribute("value").strip()


def _collect_valuations(node, found):
    """Walk the Equipment Details response, gathering every EQD_VALAMOUNT."""
    if isinstance(node, dict):
        if "EQD_VALAMOUNT" in node:
            found.append(node["EQD_VALAMOUNT"])
        for value in node.values():
            _collect_valuations(value, found)
    elif isinstance(node, list):
        for item in node:
            _collect_valuations(item, found)


def _to_number(value):
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return 0.0


_INQUIRE_JS = """
const done = arguments[arguments.length - 1];
const tokenEl = document.querySelector('input[name="__RequestVerificationToken"]');
const body = new URLSearchParams({
  strContractNo_RECContractInquiry: arguments[0],
  strInquireType_RECContractInquiry: 'MAININQ_EQUIPMENTDETAILS',
  STR_FORM_ID: '02461', STR_FUNCTION_ID: 'CR', STR_PREMIS: 'HOD',
  STR_INSTANT: 'ALCL', STR_APP_ID: '00023'
});
if (tokenEl) { body.set('__RequestVerificationToken', tokenEl.value); }
fetch(arguments[1], {
  method: 'POST', credentials: 'same-origin',
  headers: {'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
            'X-Requested-With': 'XMLHttpRequest'},
  body: body.toString()
}).then(r => r.text()).then(t => done(t)).catch(e => done('ERR:' + e));
"""


def _equipment_valuations(driver, wait, contract_no):
    """Open Equipment Details for the contract just searched and return the
    Valuation Amount of every equipment row (list of floats)."""
    try:
        btn = wait.until(EC.element_to_be_clickable((By.ID, ID_EQUIPMENT_BUTTON)))
        driver.execute_script("arguments[0].scrollIntoView(true);", btn)
        try:
            btn.click()
        except Exception:
            driver.execute_script("arguments[0].click();", btn)
        _pause(1)  # popup opens over the page - visible when running with a browser window
    except (TimeoutException, WebDriverException) as e:
        if _is_connection_error(e):
            raise
        print(f"  (Equipment Details button not clickable: {type(e).__name__})")

    valuations = []
    try:
        driver.set_script_timeout(25)
        raw = driver.execute_async_script(_INQUIRE_JS, contract_no, INQUIRE_API_URL)
        if raw and not str(raw).startswith("ERR:"):
            _collect_valuations(json.loads(raw), valuations)
        else:
            logger.warning(f"  {contract_no}: Equipment Details request failed: {raw!r}")
    except (json.JSONDecodeError, WebDriverException) as e:
        if isinstance(e, WebDriverException) and _is_connection_error(e):
            raise
        logger.warning(f"  {contract_no}: could not read Equipment Details ({type(e).__name__}: {e})")
    finally:
        try:
            close = driver.find_element(By.ID, ID_EQUIPMENT_CLOSE)
            driver.execute_script("arguments[0].click();", close)
            _pause(0.5)
        except Exception:
            pass

    return [_to_number(v) for v in valuations]


def _first_positive(valuations):
    return next((v for v in valuations if v > 0), 0.0)


def _active_contracts_for_client(driver, wait, client_code):
    """Type the client code (no Enter), click the '?' after Contract No, and read
    the contract list. Returns [(contract_no, status), ...] for every row."""
    _reset_form(driver)
    _type_into(driver, wait, ID_CLIENT_CODE_INPUT, client_code)
    print(f"  Typed client code {client_code}, clicking '?' next to Contract No...")
    icon = wait.until(EC.element_to_be_clickable((By.ID, ID_CONTRACT_HELP_ICON)))
    try:
        icon.click()
    except Exception:
        driver.execute_script("arguments[0].click();", icon)

    wait.until(EC.presence_of_element_located((By.CSS_SELECTOR, f"#{ID_HELP_TABLE_BODY} tr")))
    _pause(1)
    contracts = []
    for row in driver.find_elements(By.CSS_SELECTOR, f"#{ID_HELP_TABLE_BODY} tr"):
        cells = row.find_elements(By.TAG_NAME, "td")
        if len(cells) > HELP_COL_STATUS:
            contracts.append((cells[HELP_COL_CONTRACT].text.strip(), cells[HELP_COL_STATUS].text.strip().upper()))
    try:
        driver.execute_script("arguments[0].click();", driver.find_element(By.ID, ID_HELP_CLOSE))
        _pause(0.5)
    except Exception:
        pass
    return contracts


def find_valuation(driver, wait, contract_no, client_code_hint=None):
    """Returns {'valuation': float, 'source': str, 'client_code': str}.
    valuation is 0.0 when neither the contract nor any of its client's Active
    contracts has a Valuation Amount."""
    contract_no = str(contract_no).strip()
    print(f"-> {contract_no}: searching contract...")
    erp_client_code = _search_contract(driver, wait, contract_no)
    client_code = (client_code_hint or erp_client_code or "").strip()

    valuations = _equipment_valuations(driver, wait, contract_no)
    print(f"  Equipment Details Valuation Amount(s): {valuations or 'no records'}")
    value = _first_positive(valuations)
    if value > 0:
        print(f"  FOUND on the contract itself: {value:,.2f}")
        return {"valuation": value, "source": "contract", "client_code": client_code}

    if not client_code:
        print("  No client code available - cannot try the client's other contracts")
        return {"valuation": 0.0, "source": "none", "client_code": ""}

    print(f"  No valuation - looking for other Active contracts of client {client_code}...")
    listed = _active_contracts_for_client(driver, wait, client_code)
    print(f"  Contract list: {listed}")
    candidates = [c for c, status in listed if status == "A" and c != contract_no]
    if not candidates:
        print("  No other Active ('A') contracts for this client")

    for candidate in candidates:
        print(f"  Trying Active contract {candidate}...")
        _search_contract(driver, wait, candidate)
        valuations = _equipment_valuations(driver, wait, candidate)
        print(f"    Valuation Amount(s): {valuations or 'no records'}")
        value = _first_positive(valuations)
        if value > 0:
            print(f"  FOUND via client's Active contract {candidate}: {value:,.2f}")
            return {"valuation": value, "source": f"client-active:{candidate}", "client_code": client_code}

    print("  NOT FOUND - will be reported as a missing record")
    return {"valuation": 0.0, "source": "none", "client_code": client_code}


def _safe_find_valuation(driver, wait, contract_no, client_code_hint):
    """find_valuation with one retry after closing an ERP error dialog, and a
    failure result (rather than an exception) for any single-contract error."""
    for attempt in (1, 2):
        try:
            return find_valuation(driver, wait, contract_no, client_code_hint)
        except WebDriverException as e:
            if _is_connection_error(e):
                raise
            if attempt == 1 and _dismiss_error_dialog(driver, contract_no):
                logger.info(f"  {contract_no}: retrying after closing the ERP error dialog")
                continue
            logger.error(f"  {contract_no}: lookup failed ({type(e).__name__}: {e})")
            break
        except Exception as e:
            logger.error(f"  {contract_no}: lookup failed ({type(e).__name__}: {e})")
            break
    return {"valuation": 0.0, "source": "error", "client_code": client_code_hint or ""}


# ---------------------------------------------------------------- Excel side

def _is_blank_or_zero(value):
    return value is None or str(value).strip() in ("", "0", "0.0")


def find_na_rows(wb):
    """Rows of the LTV sheet whose Initial Valuation formula would evaluate to
    #N/A. The report is saved without calculated values, so this replays the
    formula's logic: VLOOKUP(contract, Summary!A:AY, 51) fails when the contract
    isn't on Summary (or has no valuation there), and rows whose formula carries
    an IFNA(...) fallback are already resolved."""
    ws = wb[LTV_SHEET]
    summary_valuation = {}
    for row in wb[SUMMARY_SHEET].iter_rows(min_row=3, min_col=1, max_col=SUMMARY_COL_INITIAL_VALUATION,
                                           values_only=True):
        if row[0] is not None:
            summary_valuation[str(row[0]).strip()] = row[SUMMARY_COL_INITIAL_VALUATION - 1]

    client_by_contract = {}
    if DISBURSEMENT_SHEET in wb.sheetnames:
        ws_d = wb[DISBURSEMENT_SHEET]
        for r in range(DISBURSEMENT_START_ROW, ws_d.max_row + 1):
            contract, client = ws_d.cell(row=r, column=2).value, ws_d.cell(row=r, column=1).value
            if contract is not None and client is not None:
                client_by_contract[str(contract).strip()] = _normalize_code(client)

    rows = []
    for r in range(LTV_START_ROW, ws.max_row + 1):
        contract = ws.cell(row=r, column=COL_CONTRACT).value
        formula = ws.cell(row=r, column=COL_INIT_VAL).value
        if contract is None or not isinstance(formula, str) or not formula.startswith("="):
            continue
        if "IFNA" in formula.upper():
            continue
        contract = str(contract).strip()
        if _is_blank_or_zero(summary_valuation.get(contract)):
            rows.append({"row": r, "contract": contract, "client_code": client_by_contract.get(contract)})
    return rows


def _normalize_code(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def add_to_missing_records(report_file, wb, missing_rows):
    """Write the unresolved rows, as plain values, to a 'LTV Working - Disbursement'
    sheet of Missing_Records.xlsx (created if the MT step didn't make one)."""
    path = Path(report_file).parent / "Missing_Records.xlsx"
    if path.exists():
        out = load_workbook(path)
        if LTV_SHEET in out.sheetnames:
            del out[LTV_SHEET]
        ws_out = out.create_sheet(LTV_SHEET)
    else:
        out = Workbook()
        ws_out = out.active
        ws_out.title = LTV_SHEET

    ws_out.append(LTV_HEADERS)
    ws = wb[LTV_SHEET]
    for item in missing_rows:
        r = item["row"]
        contract = str(ws.cell(row=r, column=COL_CONTRACT).value).strip()
        grant = ws.cell(row=r, column=COL_GRANT).value
        ws_out.append([
            contract,
            ws.cell(row=r, column=COL_NAME).value,
            contract[2:4],                                   # sheet formula: =MID(A,3,2)
            ws.cell(row=r, column=COL_EQUIPMENT).value,
            grant,                                           # Gross Exposure falls back to Grant Amount off-Summary
            ws.cell(row=r, column=COL_COLLATERAL).value,
            grant,
            None,                                            # Initial Valuation: none found
            None,                                            # LTV %: needs a valuation
            None,                                            # WALTV %
        ])
    out.save(path)
    return path


def run(report_file):
    wb = load_workbook(report_file)
    if LTV_SHEET not in wb.sheetnames or SUMMARY_SHEET not in wb.sheetnames:
        logger.warning(f"'{LTV_SHEET}' or '{SUMMARY_SHEET}' sheet not found - nothing to do")
        return

    na_rows = find_na_rows(wb)
    logger.info(f"{len(na_rows)} row(s) in '{LTV_SHEET}' have an unresolved (#N/A) Initial Valuation")
    if not na_rows:
        return

    driver, wait = _erp_login()
    ws = wb[LTV_SHEET]
    missing, filled = [], 0
    try:
        navigate_to_contract_inquiry(driver, wait)
        cache = {}
        for idx, item in enumerate(na_rows, start=1):
            contract = item["contract"]
            print(f"[{idx}/{len(na_rows)}]", end=" ")
            if contract not in cache:
                cache[contract] = _safe_find_valuation(driver, wait, contract, item["client_code"])
            result = cache[contract]
            if result["valuation"] > 0:
                ws.cell(row=item["row"], column=COL_INIT_VAL).value = result["valuation"]
                filled += 1
            else:
                missing.append(item)
            logger.info(f"[{idx}/{len(na_rows)}] {contract}: {result['source']} valuation={result['valuation']}")
    finally:
        driver.quit()

    if missing:
        path = add_to_missing_records(report_file, wb, missing)
        logger.info(f"{len(missing)} contract(s) without a valuation written to {path} ('{LTV_SHEET}' sheet)")

    wb.save(report_file)
    logger.info(f"Filled {filled} Initial Valuation cell(s); {len(missing)} missing. Saved: {report_file}")


def test_contracts(contract_nos):
    """Watch individual records: visible Chrome, every step printed, no Excel file
    read or written. Contracts can be 'CONTRACT' or 'CONTRACT:CLIENTCODE'."""
    driver, wait = _erp_login()
    try:
        navigate_to_contract_inquiry(driver, wait)
        for idx, entry in enumerate(contract_nos, start=1):
            contract, _, client = entry.partition(":")
            print(f"[{idx}/{len(contract_nos)}]", end=" ")
            result = _safe_find_valuation(driver, wait, contract, client.strip() or None)
            print(f"  RESULT: {result}\n")
    finally:
        input("\nDone. Press Enter to close the browser...")
        driver.quit()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Fill #N/A Initial Valuation in LTV Working - Disbursement from the ERP")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--report-file", help="Path to the Quarter Classification Loans & Advance xlsx")
    group.add_argument("--contracts", help="Comma-separated contract numbers (optionally CONTRACT:CLIENTCODE) "
                                           "to test in the browser, no file involved")
    args = parser.parse_args()

    if args.contracts:
        test_contracts([c.strip() for c in args.contracts.split(",") if c.strip()])
    else:
        run(args.report_file)
